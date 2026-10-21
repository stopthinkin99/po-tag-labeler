from __future__ import annotations

import io
import os
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import fitz  # PyMuPDF
from openpyxl import load_workbook

try:
    import pytesseract
    from PIL import Image
except ImportError:  # lets normal text PDFs keep working even before OCR deps install
    pytesseract = None
    Image = None


PO_HEADER_CANDIDATES = {
    "cust po no", "cust po number", "customer po", "customer po no",
    "customer po number", "po no", "po number",
}
TAG_HEADER_CANDIDATES = {
    "tag no", "tag number", "tag", "tag #",
}

BOLD_FONT = "Helvetica-Bold"


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\n", " ").strip().lower().split())


def normalize_po(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    text = re.sub(r"\.0$", "", text)
    return re.sub(r"[^A-Za-z0-9]", "", text).upper()


def load_po_tag_map(excel_bytes: bytes) -> dict[str, list[str]]:
    workbook = load_workbook(io.BytesIO(excel_bytes), data_only=True, read_only=True)

    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        for row_index, row in enumerate(rows[:20]):
            normalized = [normalize_header(cell) for cell in row]
            po_col = next((i for i, h in enumerate(normalized) if h in PO_HEADER_CANDIDATES), None)
            tag_col = next((i for i, h in enumerate(normalized) if h in TAG_HEADER_CANDIDATES), None)

            if po_col is None or tag_col is None:
                continue

            mapping: dict[str, list[str]] = defaultdict(list)
            for data_row in rows[row_index + 1:]:
                if po_col >= len(data_row) or tag_col >= len(data_row):
                    continue
                po = normalize_po(data_row[po_col])
                tag_value = data_row[tag_col]
                if not po or tag_value is None or str(tag_value).strip() == "":
                    continue
                if isinstance(tag_value, float) and tag_value.is_integer():
                    tag_value = int(tag_value)
                tag = str(tag_value).strip()
                if tag not in mapping[po]:
                    mapping[po].append(tag)

            if mapping:
                return dict(mapping)

    raise ValueError(
        "Could not find the Excel columns. Expected headers similar to "
        "'Cust Po No' and 'Tag NO'."
    )


def extract_po_from_page_text(text: str) -> str | None:
    """Extract the customer PO from FedEx-style text or UPS Purchase No text."""
    if not text:
        return None

    # UPS label example:
    # Purchase No.: 10000032644860-47502218
    # The customer PO is the part AFTER the final dash: 47502218.
    purchase_match = re.search(
        r"\bPURCHASE\s*NO\.?\s*:\s*([A-Z0-9]+(?:-[A-Z0-9]+)+)",
        text,
        flags=re.IGNORECASE,
    )
    if purchase_match:
        full_purchase_number = purchase_match.group(1)
        return normalize_po(full_purchase_number.rsplit("-", 1)[-1])

    patterns = [
        r"\bREF\s*:\s*([A-Z0-9][A-Z0-9\-/]*)",
        r"\bPO\s*(?:NO|NUMBER|#)?\s*:\s*([A-Z0-9][A-Z0-9\-/]*)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return normalize_po(match.group(1))
    return None


def _configure_tesseract() -> None:
    """Locate Tesseract on common Windows setups or use PATH on Linux/cloud."""
    if pytesseract is None:
        return

    env_cmd = os.getenv("TESSERACT_CMD")
    if env_cmd and Path(env_cmd).exists():
        pytesseract.pytesseract.tesseract_cmd = env_cmd
        return

    candidates = [
        Path("bundle/Tesseract-OCR/tesseract.exe"),
        Path("Tesseract-OCR/tesseract.exe"),
        Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
        Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
    ]
    for candidate in candidates:
        if candidate.exists():
            pytesseract.pytesseract.tesseract_cmd = str(candidate)
            return


def _ocr_purchase_po(page: fitz.Page) -> str | None:
    """OCR only the bottom of image-only UPS labels to read Purchase No."""
    if pytesseract is None or Image is None:
        return None

    _configure_tesseract()

    # UPS Purchase No is near the bottom, so crop to improve accuracy and speed.
    r = page.rect
    clip = fitz.Rect(r.x0, r.y0 + r.height * 0.72, r.x1, r.y1)
    pix = page.get_pixmap(matrix=fitz.Matrix(4, 4), clip=clip, alpha=False)
    image = Image.open(io.BytesIO(pix.tobytes("png")))

    try:
        ocr_text = pytesseract.image_to_string(image, config="--psm 6")
    except Exception:
        return None

    return extract_po_from_page_text(ocr_text)


@dataclass
class PageResult:
    page_number: int
    po_number: str | None
    tags: list[str]
    status: str
    source: str = ""


def _overlap_area(a: fitz.Rect, b: fitz.Rect) -> float:
    inter = a & b
    if inter.is_empty:
        return 0.0
    return max(0.0, inter.width) * max(0.0, inter.height)


def _find_safe_top_box(page: fitz.Page, box_w: float, box_h: float) -> tuple[float, float]:
    page_rect = page.rect

    occupied = []
    for word in page.get_text("words"):
        x0, y0, x1, y1 = word[:4]
        occupied.append(fitz.Rect(x0 - 1, y0 - 1, x1 + 1, y1 + 1))

    x_centers = [0.88, 0.50, 0.75, 0.25, 0.62, 0.38]
    y_values = [4, 12, 20, 28, 36, 44, 52, 60, 68, 76, 84]
    best: tuple[float, float, float] | None = None

    for y0 in y_values:
        for center in x_centers:
            x0 = page_rect.width * center - box_w / 2
            x0 = min(max(4, x0), page_rect.width - box_w - 4)
            candidate = fitz.Rect(x0, y0, x0 + box_w, y0 + box_h)
            overlap = sum(_overlap_area(candidate, rect) for rect in occupied)
            score = overlap * 1000 + y0 + abs(center - 0.88) * 8
            if best is None or score < best[0]:
                best = (score, x0, y0)
            if overlap == 0:
                return x0, y0

    assert best is not None
    return best[1], best[2]


def _already_has_top_tag(page: fitz.Page, tags: Iterable[str]) -> bool:
    """Avoid adding a duplicate when reprocessing a PDF already tagged by this app."""
    tag_values = {str(t).strip() for t in tags}
    if not tag_values:
        return False

    r = page.rect
    for word in page.get_text("words"):
        x0, y0, x1, y1, text = word[:5]
        if (
            text.strip() in tag_values
            and x0 > r.width * 0.78
            and y0 < r.height * 0.12
        ):
            return True
    return False


def _remove_old_tag_box(page: fitz.Page) -> None:
    """Cover an older large 'TAG NO' stamp if an already-processed PDF is reused."""
    matches = []
    for needle in ("TAG NO:", "TAG NO"):
        try:
            matches.extend(page.search_for(needle))
        except Exception:
            pass

    candidates = [r for r in matches if r.y0 < 120]
    if not candidates:
        return

    rect = candidates[0]
    for r in candidates[1:]:
        rect |= r

    cleanup = fitz.Rect(
        max(page.rect.x0, rect.x0 - 85),
        max(page.rect.y0, rect.y0 - 18),
        min(page.rect.x1, rect.x1 + 85),
        min(page.rect.y1, rect.y1 + 20),
    )
    page.draw_rect(cleanup, color=(1, 1, 1), fill=(1, 1, 1), width=0, overlay=True)


def _stamp_tag(
    page: fitz.Page,
    tags: Iterable[str],
    position: str = "auto-safe-top",
    font_size: int = 10,
) -> None:
    tag_text = ", ".join(tags)
    page_rect = page.rect
    actual_size = max(7, min(font_size, 16))
    text_width = fitz.get_text_length(tag_text, fontname=BOLD_FONT, fontsize=actual_size)

    circle_padding = max(5, actual_size * 0.55)
    diameter = max(text_width + circle_padding * 2, actual_size * 2.4)
    margin = 6

    if position == "ups-safe":
        # UPS image labels have no text layer, so reserve the open right side
        # of the SHIP TO area instead of guessing from PDF text coordinates.
        x0 = page_rect.width - diameter - 8
        y0 = min(max(52, page_rect.height * 0.16), page_rect.height - diameter - 8)
    elif position == "auto-safe-top":
        x0, y0 = _find_safe_top_box(page, diameter, diameter)
    else:
        if position == "top-left":
            x0 = margin
        elif position == "top-center":
            x0 = (page_rect.width - diameter) / 2
        else:
            x0 = page_rect.width - diameter - margin
        y0 = margin

    center_x = x0 + diameter / 2
    center_y = y0 + diameter / 2
    radius = diameter / 2

    page.draw_circle(
        (center_x, center_y),
        radius,
        color=(0, 0, 0),
        width=1.3,
        fill=None,
        overlay=True,
    )

    text_x = center_x - text_width / 2
    baseline_y = center_y + actual_size * 0.36
    page.insert_text(
        (text_x, baseline_y),
        tag_text,
        fontsize=actual_size,
        fontname=BOLD_FONT,
        color=(0, 0, 0),
        overlay=True,
    )


def process_pdf(
    pdf_bytes: bytes,
    mapping: dict[str, list[str]],
    position: str = "auto-safe-top",
    font_size: int = 10,
    remove_old_tag_box: bool = True,
) -> tuple[bytes, list[PageResult]]:
    document = fitz.open(stream=pdf_bytes, filetype="pdf")
    results: list[PageResult] = []

    for index, page in enumerate(document):
        page_text = page.get_text("text") or ""
        po = extract_po_from_page_text(page_text)
        source = "PDF text"

        # Image-only UPS labels have no selectable text. OCR the bottom strip,
        # find "Purchase No.", and use the number after the final dash.
        if not po:
            po = _ocr_purchase_po(page)
            if po:
                source = "UPS Purchase No (OCR)"

        if not po:
            results.append(PageResult(index + 1, None, [], "PO/REF not found", ""))
            continue

        tags = mapping.get(po, [])
        if not tags:
            results.append(PageResult(index + 1, po, [], "No Excel match", source))
            continue

        if remove_old_tag_box:
            _remove_old_tag_box(page)

        if not _already_has_top_tag(page, tags):
            stamp_position = "ups-safe" if source == "UPS Purchase No (OCR)" and position == "auto-safe-top" else position
            _stamp_tag(page, tags, position=stamp_position, font_size=font_size)
            status = "Matched"
        else:
            status = "Matched (tag already present)"

        results.append(PageResult(index + 1, po, tags, status, source))

    output = document.tobytes(garbage=4, deflate=True)
    document.close()
    return output, results
